"""Create slide.pptx, the original PowerPoint input for the README showcase image.

Requires python-pptx (pip install python-pptx). Run from this folder:

    python make-slide.py
"""

from pptx import Presentation
from pptx.chart.data import CategoryChartData
from pptx.dml.color import RGBColor
from pptx.enum.chart import XL_CHART_TYPE
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import PP_ALIGN
from pptx.util import Inches, Pt

prs = Presentation()
prs.slide_width = Inches(10)
prs.slide_height = Inches(7.5)
slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank layout


def add_box(shape_type, left, top, width, height, color, text, size):
    """Add a filled shape with centered white text."""
    shape = slide.shapes.add_shape(
        shape_type, Inches(left), Inches(top), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = RGBColor.from_string(color)
    shape.line.fill.background()
    frame = shape.text_frame
    frame.text = text
    frame.paragraphs[0].alignment = PP_ALIGN.CENTER
    run = frame.paragraphs[0].runs[0]
    run.font.size = Pt(size)
    run.font.bold = True
    run.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)
    return shape


add_box(MSO_SHAPE.ROUNDED_RECTANGLE, 0.5, 0.4, 9.0, 1.1, "2563EB", "Quarterly Review", 36)

chart_data = CategoryChartData()
chart_data.categories = ["Q1", "Q2", "Q3", "Q4"]
chart_data.add_series("North", (42, 55, 61, 78))
chart_data.add_series("South", (30, 38, 52, 60))
chart = slide.shapes.add_chart(
    XL_CHART_TYPE.COLUMN_CLUSTERED, Inches(0.5), Inches(1.8), Inches(5.6), Inches(5.0), chart_data
).chart
chart.has_legend = True
chart.has_title = False
axis = chart.value_axis
axis.minimum_scale = 0
axis.maximum_scale = 80
axis.major_unit = 20

add_box(MSO_SHAPE.OVAL, 6.5, 2.0, 3.0, 1.6, "7C3AED", "+38%", 32)
add_box(MSO_SHAPE.ROUNDED_RECTANGLE, 6.5, 4.0, 3.0, 1.2, "059669", "On target", 24)
add_box(MSO_SHAPE.PENTAGON, 6.5, 5.5, 3.0, 1.2, "D97706", "Next steps", 24)

prs.save("slide.pptx")
